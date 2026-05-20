"""
File parser module — Extract text from various document formats.
Supports: .txt, .pdf, .docx
"""
import os
import logging
from pathlib import Path

logger = logging.getLogger(__name__)

# Optional imports for PDF and DOCX
try:
    import pypdf
    HAS_PYPDF = True
except ImportError:
    HAS_PYPDF = False
    logger.warning("pypdf not installed. PDF support disabled. Install with: pip install pypdf")

try:
    from docx import Document
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False
    logger.warning("python-docx not installed. DOCX support disabled. Install with: pip install python-docx")


def parse_txt(file_path: str) -> str:
    """Extract text from a .txt file."""
    try:
        with open(file_path, 'r', encoding='utf-8') as f:
            text = f.read()
        logger.info(f"Parsed TXT: {file_path} ({len(text)} chars)")
        return text
    except Exception as e:
        logger.error(f"Error parsing TXT {file_path}: {e}")
        raise


def parse_pdf(file_path: str) -> str:
    """Extract text from a .pdf file using pypdf."""
    if not HAS_PYPDF:
        raise ImportError("pypdf is required for PDF support. Install with: pip install pypdf")
    
    try:
        text_parts = []
        with open(file_path, 'rb') as f:
            reader = pypdf.PdfReader(f)
            for page_num, page in enumerate(reader.pages):
                text = page.extract_text()
                if text:
                    text_parts.append(text)
        
        text = "\n".join(text_parts)
        logger.info(f"Parsed PDF: {file_path} ({len(text)} chars from {len(reader.pages)} pages)")
        return text
    except Exception as e:
        logger.error(f"Error parsing PDF {file_path}: {e}")
        raise


def parse_docx(file_path: str) -> str:
    """Extract text from a .docx file using python-docx."""
    if not HAS_DOCX:
        raise ImportError("python-docx is required for DOCX support. Install with: pip install python-docx")
    
    try:
        doc = Document(file_path)
        text_parts = [para.text for para in doc.paragraphs if para.text.strip()]
        text = "\n".join(text_parts)
        logger.info(f"Parsed DOCX: {file_path} ({len(text)} chars from {len(doc.paragraphs)} paragraphs)")
        return text
    except Exception as e:
        logger.error(f"Error parsing DOCX {file_path}: {e}")
        raise


def parse_file(file_path: str) -> str:
    """
    Parse a file based on its extension.
    Returns extracted text.
    Raises ValueError for unsupported file types.
    """
    file_path = str(file_path)
    ext = Path(file_path).suffix.lower()
    
    if ext == '.txt':
        return parse_txt(file_path)
    elif ext == '.pdf':
        return parse_pdf(file_path)
    elif ext == '.docx':
        return parse_docx(file_path)
    else:
        raise ValueError(f"Unsupported file type: {ext}. Supported: .txt, .pdf, .docx")


def is_supported_file(filename: str) -> bool:
    """Check if a file type is supported."""
    ext = Path(filename).suffix.lower()
    return ext in ['.txt', '.pdf', '.docx']


def validate_file(file_path: str, max_size_mb: int = 50) -> tuple:
    """
    Validate file before parsing.
    Returns (is_valid, error_message).
    """
    file_path = str(file_path)
    
    if not os.path.exists(file_path):
        return False, "File does not exist"
    
    if not is_supported_file(file_path):
        return False, f"Unsupported file type. Supported: .txt, .pdf, .docx"
    
    file_size_mb = os.path.getsize(file_path) / (1024 * 1024)
    if file_size_mb > max_size_mb:
        return False, f"File too large ({file_size_mb:.1f}MB > {max_size_mb}MB)"
    
    return True, ""
